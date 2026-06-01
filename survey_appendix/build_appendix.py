#!/usr/bin/env python3
"""
Order Process Survey -- Appendix generator.

Builds a PDF appendix of individual survey responses, mirroring the format of
the "AI Usage Survey" appendix (navy/steel card per respondent).

Usage:
    python3 build_appendix.py --limit 10 --out /tmp/appendix_sample.pdf
    python3 build_appendix.py --out /tmp/appendix_full.pdf            # everything
    python3 build_appendix.py --named-only --out appendix_named.pdf   # skip anonymous
"""
import argparse
import html
import re
import openpyxl
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether,
)

XLSX = "/root/.claude/uploads/15a9b835-0239-43a5-816e-283ad0ca0bc2/7c8119a2-Survey_Responses.xlsx"

# ---- palette (sampled from the AI Usage Survey appendix) --------------------
NAVY        = colors.HexColor("#1f2d4d")   # "Marketed" header bar
STEEL       = colors.HexColor("#4a6791")   # "Not marketed / unknown" header bar
INK         = colors.HexColor("#1f2d4d")   # bold labels / title
GREY_TEXT   = colors.HexColor("#6b7280")   # subtitle
BAND        = colors.HexColor("#eef1f6")   # info-band background
BORDER      = colors.HexColor("#c8d0de")   # card border
HASH_COLOR  = colors.HexColor("#b8c4dc")   # the "#N" on the header bar
BODY_INK    = colors.HexColor("#2b3440")

# ---- column indexes (1-based, from the Raw Data tab) ------------------------
C = dict(
    q1=18, q1_text=19, q2=20, q2_text=21, q3=22, q3_text=23,
    q4=24, q4_text=25, q5=26, q6=27,
    q7b_start=37, q7b_other=45,         # "would use on McMaster" grid (8 cols) + other text
    q8=46, role=47, role_text=48, q10=49,
    name=50, company=51, email=52, phone=53,
    marketed=59, cmf=60, mmf=61,
)

FEATURES = [
    "Add teammates to your account",
    "Send orders for approval",
    "Set roles (requester, buyer, approver)",
    "Attach purchase orders",
    "Assign account / job / cost-center codes",
    "Share order status with others",
    "Connect with your procurement system",
    "Other",
]

MAX_IDS = 5   # cap CMF/MMF numbers shown per card (after dedup)


# ----------------------------------------------------------------------------
def clean(v):
    if v is None:
        return ""
    s = str(v).strip()
    if s.upper() == "NULL":
        return ""
    return s


def esc(v):
    return html.escape(clean(v))


def split_multi(v, other_text="", other_label="Other"):
    """Split a comma-joined multi-select; swap the 'Other (please ...)' token
    for the respondent's written-in text."""
    s = clean(v)
    if not s:
        return []
    # Qualtrics joins selected options with a comma+no-space; commas *inside*
    # an option label are always followed by a space. Split only on the former.
    out = []
    for tok in re.split(r",(?!\s)", s):
        tok = tok.strip()
        if not tok:
            continue
        if tok.lower().startswith("other"):
            ot = clean(other_text)
            out.append(f'{other_label}: "{ot}"' if ot else other_label)
        else:
            out.append(tok)
    return out


def ids_field(v):
    """Tidy an account-id list: drop NULL/0/blank, de-duplicate (MMF often
    repeats the same id many times), and cap the display (the source lists can
    run to 100+ entries, but only a few are useful on the card)."""
    s = clean(v)
    if not s:
        return ""
    seen = []
    for p in s.split(","):
        p = p.strip()
        if p and p.upper() != "NULL" and p != "0" and p not in seen:
            seen.append(p)
    if not seen:
        return ""
    if len(seen) > MAX_IDS:
        return ", ".join(seen[:MAX_IDS]) + f" (+{len(seen) - MAX_IDS} more)"
    return ", ".join(seen)


def marketed_label(v):
    s = clean(v).lower()
    if s == "yes":
        return "Marketed", NAVY
    if s == "no":
        return "Not Marketed", STEEL
    return "Unknown", STEEL


# ---- paragraph styles -------------------------------------------------------
TITLE = ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=24,
                       textColor=INK, alignment=1, spaceAfter=6, leading=28)
SUBTITLE = ParagraphStyle("subtitle", fontName="Helvetica", fontSize=11.5,
                          textColor=GREY_TEXT, alignment=1, leading=15)
HEAD = ParagraphStyle("head", fontName="Helvetica-Bold", fontSize=9.5,
                      textColor=colors.white, leading=12)
HEADNUM = ParagraphStyle("headnum", fontName="Helvetica-Bold", fontSize=9.5,
                         textColor=HASH_COLOR, alignment=2, leading=12)
INFO = ParagraphStyle("info", fontName="Helvetica", fontSize=8.6,
                      textColor=BODY_INK, leading=13)
BLABEL = ParagraphStyle("blabel", fontName="Helvetica-Bold", fontSize=8.8,
                        textColor=INK, leading=12, spaceBefore=4, spaceAfter=1)
BTEXT = ParagraphStyle("btext", fontName="Helvetica", fontSize=8.6,
                       textColor=BODY_INK, leading=12.5)
BBULLET = ParagraphStyle("bbullet", fontName="Helvetica", fontSize=8.6,
                         textColor=BODY_INK, leading=12.5, leftIndent=10,
                         bulletIndent=0)


def kv(label, value):
    return f'<b>{label}:</b> {value}' if value else ""


def build_card(row, n, content_w):
    g = lambda key: clean(row.get(key))
    name = g("name")
    company = g("company")
    role = g("role")
    if role.lower().startswith("other"):
        rt = g("role_text")
        role = f"Other: {rt}" if rt else "Other"

    mlabel, mcolor = marketed_label(row.get("marketed"))

    # ---- header bar ----
    title_bits = [name or "Anonymous"]
    if company:
        title_bits.append(company)
    if role:
        title_bits.append(role)
    title_bits.append(mlabel)
    header_left = Paragraph(esc(" | ".join(title_bits)), HEAD)
    header_right = Paragraph(f"#{n}", HEADNUM)
    header = Table([[header_left, header_right]],
                   colWidths=[content_w * 0.82, content_w * 0.18])
    header.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), mcolor),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (0, 0), 9),
        ("RIGHTPADDING", (-1, 0), (-1, 0), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))

    # ---- info band (2 columns) ----
    email = esc(row.get("email"))
    q1 = g("q1")
    if "depends" in q1.lower():
        t = g("q1_text")
        q1 = "It depends" + (f" — {t}" if t else "")
    q2 = g("q2")
    if "sometimes" in q2.lower():
        t = g("q2_text")
        q2 = "Sometimes" + (f" — {t}" if t else "")
    left_lines = [ln for ln in [
        kv("Email", email),
        kv("Buying process", esc(q1)),
        kv("Approval needed", esc(q2)),
    ] if ln]
    right_lines = [ln for ln in [
        kv("CMF", esc(ids_field(row.get("cmf")))),
        kv("MMF", esc(ids_field(row.get("mmf")))),
        kv("Re-enters details", esc(g("q5"))),
    ] if ln]
    info = Table([[Paragraph("<br/>".join(left_lines) or "&nbsp;", INFO),
                   Paragraph("<br/>".join(right_lines) or "&nbsp;", INFO)]],
                 colWidths=[content_w * 0.56, content_w * 0.44])
    info.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), BAND),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (0, 0), 9),
        ("LEFTPADDING", (1, 0), (1, 0), 6),
        ("RIGHTPADDING", (-1, 0), (-1, 0), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))

    # ---- body ----
    body = []
    handoff = split_multi(row.get("q3"), row.get("q3_text"))
    if handoff:
        body.append(Paragraph(f"<b>Handed off via:</b> {esc('; '.join(handoff))}", BTEXT))

    wanted = []
    for i, feat in enumerate(FEATURES):
        val = clean(row.get(f"q7b_{i}")).lower()
        if "regularly" in val:
            tag = "would use regularly"
        elif "occasionally" in val:
            tag = "would use occasionally"
        else:
            continue
        label = feat
        if feat == "Other":
            ot = clean(row.get("q7b_other"))
            label = f'Other: "{ot}"' if ot else "Other"
        wanted.append(f"{label} <font color='#6b7280'>({tag})</font>")
    if wanted:
        body.append(Paragraph("Features they'd use on mcmaster.com:", BLABEL))
        for w in wanted:
            body.append(Paragraph(w, BBULLET, bulletText="•"))

    q8 = esc(row.get("q8"))
    if q8:
        body.append(Paragraph("Most frustrating part of their buying process:", BLABEL))
        body.append(Paragraph(q8, BTEXT))

    q10 = esc(row.get("q10"))
    if q10:
        body.append(Paragraph("Anything else they shared:", BLABEL))
        body.append(Paragraph(q10, BTEXT))

    if not body:
        body.append(Paragraph("<i>No further detail provided.</i>", BTEXT))

    body_tbl = Table([[body]], colWidths=[content_w])
    body_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
        ("LEFTPADDING", (0, 0), (-1, -1), 9),
        ("RIGHTPADDING", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))

    # ---- wrap header + info + body in one bordered card ----
    card = Table([[header], [info], [body_tbl]], colWidths=[content_w])
    card.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ("BOX", (0, 0), (-1, -1), 0.6, BORDER),
        ("LINEBELOW", (0, 1), (-1, 1), 0.5, BORDER),
    ]))
    return card


def load_rows():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Raw Data"]
    rows = []
    for r in range(3, ws.max_row + 1):
        row = {k: ws.cell(row=r, column=col).value for k, col in C.items()
               if k not in ("q7b_start",)}
        # 8-wide "would use on McMaster" grid
        for i in range(8):
            row[f"q7b_{i}"] = ws.cell(row=r, column=C["q7b_start"] + i).value
        # keep only rows that actually have answer content
        if any(clean(row.get(k)) for k in ("q1", "q2", "q8", "q10", "name")):
            rows.append(row)
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=0, help="max cards (0 = all)")
    ap.add_argument("--named-only", action="store_true")
    ap.add_argument("--out", default="/tmp/appendix_sample.pdf")
    args = ap.parse_args()

    rows = load_rows()
    if args.named_only:
        rows = [r for r in rows if clean(r.get("name"))]

    doc = SimpleDocTemplate(
        args.out, pagesize=letter,
        leftMargin=0.85 * inch, rightMargin=0.85 * inch,
        topMargin=0.7 * inch, bottomMargin=0.7 * inch,
        title="Appendix - Survey Results",
    )
    content_w = doc.width

    story = [
        Paragraph("Appendix - Survey Results", TITLE),
        Spacer(1, 4),
        Paragraph(f"{len(rows):,} responses | May 13, 2026 to May 15, 2026", SUBTITLE),
        Spacer(1, 18),
    ]

    shown = rows if args.limit == 0 else rows[:args.limit]
    for n, row in enumerate(shown, start=1):
        story.append(KeepTogether(build_card(row, n, content_w)))
        story.append(Spacer(1, 12))

    doc.build(story)
    print(f"Wrote {args.out}  ({len(shown)} cards of {len(rows)} total responses)")


if __name__ == "__main__":
    main()
