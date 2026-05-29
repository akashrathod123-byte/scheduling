"""Build SQL to check if scanned part was later ordered by the triangulated CMF."""
import re
from openpyxl import load_workbook

ENRICHED = '/home/user/scheduling/QR_Code_Data_enriched.xlsx'
OUT_SQL  = '/home/user/scheduling/check_scan_to_order.sql'

ws = load_workbook(ENRICHED)['Cleaned Up']
hdr = [c.value for c in ws[1]]
idx = {h: i for i, h in enumerate(hdr)}

def cleanv(v):
    if v is None: return None
    s = str(v).strip()
    return s if s and s.upper() != 'NULL' else None

triples = []  # (visit_id, scan_date_int, part_upper, cmf_id)
for row in ws.iter_rows(min_row=2, values_only=True):
    vid  = cleanv(row[idx['visit_id']])
    part = cleanv(row[idx['friendly_url_txt']])
    ts   = row[idx['search_start_user_action_server_ts']]
    cmfs = cleanv(row[idx.get('cmf_ids (top 10, CC-filtered)', idx.get('cmf_ids (top 10)', -1))])
    if not vid or not part or ts is None or not cmfs:
        continue
    # part: strip to just the part number, uppercase
    part_upper = re.sub(r'[^A-Z0-9]', '', part.upper())
    if not part_upper:
        continue
    # scan_date as yyyymmdd int
    try:
        scan_date = int(ts.strftime('%Y%m%d'))
    except Exception:
        continue
    for cmf in cmfs.split('|'):
        cmf = cmf.strip()
        if cmf.isdigit():
            triples.append((vid, scan_date, part_upper, cmf))

print(f"Generated {len(triples):,} (visit, cmf, part) rows")

# Build SQL
values = ",\n        ".join(
    f"('{v}', {d}, '{p}', {c})" for (v, d, p, c) in triples
)
sql = f"""-- Did the triangulated CMF later order the scanned part?
WITH scans (visit_id, scan_date_int, part_upper, cmf_id) AS (
    VALUES
        {values}
)
SELECT
    s.visit_id,
    s.part_upper,
    s.cmf_id,
    s.scan_date_int                         AS scan_date,
    COUNT(DISTINCT os.wo_id)                AS orders_after_scan,
    MIN(os.entry_date)                      AS first_order_after_scan,
    DATEDIFF(day,
             CAST(CAST(s.scan_date_int AS VARCHAR(8)) AS DATE),
             CAST(CAST(MIN(os.entry_date) AS VARCHAR(8)) AS DATE)
            )                               AS days_to_first_order,
    SUM(CAST(os.sell_value AS DECIMAL(18,2))) AS total_value_after_scan
FROM scans s
LEFT JOIN secure.order_source os
    ON  os.ic1_part_number = s.part_upper
    AND (os.ship_to_cmf = s.cmf_id OR os.bill_to_cmf = s.cmf_id)
    AND os.entry_date    >= s.scan_date_int
    AND os.cancel_indicator = 'N'
    AND os.hk_stat_cd       = 'ACTIVE'
GROUP BY s.visit_id, s.part_upper, s.cmf_id, s.scan_date_int
ORDER BY orders_after_scan DESC, s.visit_id;
"""

with open(OUT_SQL, 'w') as f:
    f.write(sql)
print(f"Wrote {OUT_SQL}")
print(f"\n--- preview (first 1.5K chars) ---\n{sql[:1500]}")
