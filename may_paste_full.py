"""Reconcile pasted May SQL output against current Cleaned Up tab."""
from openpyxl import load_workbook

# (visit_id, ts_str, ip, part, user_action_engagement_count, platform, browser, special)
# Extracted from the user's paste. Internal IP 50.233.58.1 filtered downstream.
PASTE = [
    ('b053a4e442e144078effae1b6213a83a', '2026-05-30 10:30:51', '207.212.6.42',    '91578a863',  1, 'mobile_website', 'ios safari', ''),
    ('105e08cf961a422b94555b50e147c208', '2026-05-30 10:31:03', '107.115.224.5',   '91578a863',  1, 'mobile_website', 'ios safari', ''),
    ('98c0df6201b7413297956884df8795ec', '2026-05-28 15:36:14', '47.206.56.119',   '91400a108',  5, 'mcmaster_com',   'ios safari ipad', ''),
    ('c53032d3d55d43c2a6b0e7ec73da6cff', '2026-05-28 15:38:43', '47.206.56.119',   '93085a194',  4, 'mcmaster_com',   'ios safari ipad', ''),
    ('65e36aa6e2d34873b2ff436fa0f8a897', '2026-05-28 10:15:36', '71.85.139.66',    '91255a196',  6, 'mobile_website', 'ios safari', ''),
    ('fc628f8393ab4c8da41aa2ac4dd60d76', '2026-05-28 07:09:10', '50.204.105.150',  '92200a196',  1, 'mobile_website', 'ios safari', ''),
    ('35c5e7dbd12f4ddaa0339ba4e9cbc3ca', '2026-05-28 00:58:42', '160.249.11.33',   '91812a215',  5, 'mobile_website', 'ios safari', ''),
    ('dfa1d3163f1a47aba2b227ef96f3590c', '2026-05-28 05:22:14', '166.199.114.9',   '90185a589',  3, 'mobile_website', 'android chrome', ''),
    ('6f5cb83c90f846e2ab747d2f33ef320c', '2026-05-27 09:56:00', '174.242.37.212',  '92188a145',  1, 'mobile_website', 'ios safari', ''),
    ('b72581cb9cc64938b426659751da2ba8', '2026-05-26 06:55:37', '64.118.18.137',   '97467a117',  1, 'mobile_website', 'android chrome', ''),
    ('70544a04d68a42aa92142f53dbbbff38', '2026-05-30 07:05:53', '50.151.209.1',    '92865a133',  1, 'mobile_website', 'android chrome', ''),
    ('cc9f4beeddd74b0f98fefd5199a1952b', '2026-05-29 15:05:19', '50.233.58.1',     '93298a155',  1, 'mobile_website', 'ios safari', 'INTERNAL'),
    ('90c345a7989646549ca52509fc7ee881', '2026-05-25 09:57:15', '70.25.244.66',    '3528t442',   1, 'mobile_website', 'android chrome', ''),
    ('f9c694b518fd44db872b3005d5e07b7d', '2026-05-22 07:09:29', '172.59.222.65',   '91400a196',  1, 'mobile_website', 'android chrome', ''),
    ('966dcc1b35f74130884072e86de41276', '2026-05-22 09:08:07', '200.68.164.39',   '91102a770',  1, 'mobile_website', 'android chrome', ''),
    ('9330f4a6d8274825b55b0552987bd215', '2026-05-25 06:34:08', '50.233.58.1',     '93135a146',  1, 'mobile_website', 'ios safari', 'INTERNAL'),
    ('bc169c13877541268ec827849103bd52', '2026-05-23 11:24:56', '146.75.146.171',  '92865a546',  1, 'mobile_website', 'ios safari', ''),
    ('9b64a7e826fe4d05829a47569a99fdd6', '2026-05-21 18:46:10', '138.51.72.73',    '91251a451',  6, 'mobile_website', 'ios safari', ''),
    ('fe180e481ab94c7e82b6e76394627ecf', '2026-05-21 14:47:28', '199.7.157.19',    '6266n15',    1, 'mobile_website', 'android chrome', ''),
    ('d28015f6a1ca43f6b7c60070e8345420', '2026-05-21 13:54:40', '216.236.85.10',   '8087k12',    1, 'mobile_website', 'ios safari', ''),
    ('8e0779f3f0394955b14c86b761f5ef17', '2026-05-20 15:08:06', '108.62.1.10',     '91732a770',  1, 'mobile_website', 'android chrome pixel', ''),
    ('1e65ebdc3a6f4533870bc7b3284551e5', '2026-05-20 10:39:47', '187.230.66.59',   '96246a200',  1, 'mobile_website', 'android chrome', ''),
    ('d4d32d277f4c4b72b8895d6ca38ea339', '2026-05-20 10:40:44', '23.121.6.225',    '91732a649',  1, 'mobile_website', 'android chrome', ''),
    ('08709b67cedc42be9fce31b4bb961195', '2026-05-20 07:50:13', '69.24.224.134',   '91358a503',  1, 'mobile_website', 'ios chrome', ''),
    ('5ce14dc59f08483da64a39fc750b18fe', '2026-05-20 07:36:12', '142.126.102.31',  '91732a264',  1, 'mobile_website', 'android chrome', ''),
    ('dd003b0e90374282918dc9bf96a813a0', '2026-05-21 10:11:47', '96.83.14.198',    '91280a712',  4, 'mobile_website', 'android chrome', ''),
    ('6cd8ecb67ca040bb8c1cdd4dd5863287', '2026-05-21 07:40:31', '172.59.210.82',   '98952a055',  4, 'mobile_website', 'ios safari', ''),
    ('498d6d26712f428f9dd5b87cf03e0ecb', '2026-05-20 09:49:38', None,              '97802a316',  1, None,             None, ''),
    ('498b26a9a4a1433db6284473ea6a25b9', '2026-05-20 01:01:05', '152.57.19.82',    '91000a469',  1, 'mobile_website', 'ios chrome', ''),
    ('b36307fed5124cb2a8e8dd499e8f41e5', '2026-05-20 01:07:15', '152.57.8.248',    '91000a469',  1, 'mobile_website', 'android chrome', ''),
    ('b36307fed5124cb2a8e8dd499e8f41e5', '2026-05-20 01:35:33', '152.57.8.248',    '91000a469',  1, 'mobile_website', 'android chrome', ''),
    ('cc4ac6dc342741fd8a207c24ada8af62', '2026-05-19 12:15:40', '172.59.104.107',  '51525k531',  1, 'mobile_website', 'ios safari', ''),
    ('1623e7e4ec2c46b9a9515948d9ae908b', '2026-05-19 11:20:46', '166.199.31.60',   '97111a416',  9, 'mobile_website', 'ios safari', ''),
    ('319525edc950483abe31b5935f22ee52', '2026-05-18 17:14:37', '24.182.41.162',   '92423a223',  1, 'mobile_website', 'ios safari', ''),
    ('f6e0720323b64abeb0cca9fd83994f33', '2026-05-18 16:19:56', '172.56.233.28',   '91251a242',  4, 'mobile_website', 'ios chrome', ''),
    ('08709b67cedc42be9fce31b4bb961195', '2026-05-20 07:38:09', '69.24.224.134',   '91358a503',  1, 'mobile_website', 'ios chrome', ''),
    ('9aec81fbe9c04a3daae82aee6cdae7b7', '2026-05-19 15:08:02', '174.238.96.72',   '91400a196',  1, 'mobile_website', 'android chrome', ''),
    ('0f753f478fac46d48c46d065912623cd', '2026-05-19 19:01:51', '24.159.245.95',   '98778a543',  1, 'mobile_website', 'ios safari', ''),
    ('6c3cba5987264d438453491ff889f0ca', '2026-05-19 09:26:45', '70.166.98.13',    '91732a286',  1, 'mobile_website', 'android chrome', ''),
    ('1924cd551e78437f8f73b70d3cb927b9', '2026-05-15 22:22:17', '174.206.165.113', '92865a722',  1, 'mobile_website', 'ios safari', ''),
    ('0a4128b9dba6436283b35a8d47fa07ea', '2026-05-18 13:43:54', '173.165.145.173', '91074a169',  1, 'mobile_website', 'ios chrome', ''),
    ('eed5d4e9a7214a8bbcb2093e1a55d4f6', '2026-05-18 15:05:55', '4.2.163.82',      '97467a121',  1, 'mobile_website', 'ios safari', ''),
    ('29bfe5020c524cceb2521753245f37dc', '2026-05-17 16:53:31', '166.205.97.10',   '1095k11',    1, 'mobile_website', 'ios safari', ''),
    ('d49336dff4cd4030ae3e78e6e0c46eff', '2026-05-15 17:43:52', '206.214.239.8',   '90374a111',  2, 'mcmaster_com',   'windows chrome', ''),
    ('03442b14ff6f4bc78a2bf434fed76e76', '2026-05-14 16:02:05', '174.195.194.244', '97525a665',  1, 'mobile_website', 'ios safari', ''),
    ('c34f2d6678a440a3adcc3443e2fa74ff', '2026-05-15 17:05:52', '172.59.0.47',     '97111a212',  4, 'mobile_website', 'ios safari', ''),
    ('ce51129809a841f6951a421a31bee38d', '2026-05-15 13:17:04', '206.214.239.8',   '90374a111',  2, 'mcmaster_com',   'windows chrome', ''),
    ('07f96327fb1f4e8990d507a3e0e9a5d4', '2026-05-15 05:21:10', '166.181.83.100',  '97802a323',  1, 'mobile_website', 'ios safari', ''),
    ('1a23d70be3fa40a286722b1080acacd0', '2026-05-15 08:51:35', '174.251.64.45',   '91251a245',  1, 'mobile_website', 'ios safari', ''),
    ('6e77d3a6918e42b2845a410daa567c3b', '2026-05-15 07:04:32', '174.228.226.96',  '92188a145',  7, 'mobile_website', 'ios safari', 'HENKELS'),
    ('a741feebe2854a3684a5c97dcf7216ee', '2026-05-14 14:45:22', '172.58.127.147',  '95105a159',  4, 'mobile_website', 'ios chrome', ''),
    ('7cb2b944e0ae496cac80033bde08dd17', '2026-05-13 16:08:19', '107.193.86.49',   '7113k215',   1, 'mobile_website', 'ios firefox', ''),
    ('c6e79a2e247341489edd1beb933cfd4e', '2026-05-13 13:32:30', '172.56.66.81',    '97802a316', 19, 'mobile_website', 'android chrome', ''),
    ('fd493c6653b643c19710ac23de7a8666', '2026-05-12 09:17:45', '172.58.146.85',   '90578a563',  1, 'mobile_website', 'ios safari', ''),
    ('b4592ff2e8df4e749c9018120ac7b383', '2026-05-11 21:00:41', '150.252.244.177', '92018a650',  1, 'mcmaster_com',   'windows chrome', ''),
    ('3d120cafd0ff4a3f83d5d3f6c9d7698c', '2026-05-14 05:17:57', '50.233.58.1',     '5661k55',    1, 'mobile_website', 'ios safari', 'INTERNAL'),
    ('efb14cd5d432435caf6a69c0c6d9929d', '2026-05-13 08:56:53', '172.56.140.174',  '92970a224',  1, 'mobile_website', 'ios safari', ''),
    ('1db85c2837df44e18b5d17267e62e340', '2026-05-12 10:36:07', '68.70.150.230',   '94459a130',  6, 'android_phone_app', 'android sm-s918w', ''),
    ('f9588f11f801474bb8655a14be27effa', '2026-05-11 18:55:48', '50.233.58.1',     '91287a258',  1, 'mobile_website', 'ios safari', 'INTERNAL'),
    ('A578F7659FC248C6B12465F40A20B839', '2026-05-11 18:56:10', '50.233.58.1',     '91287a258',  1, 'iphone_app',     'ios iphone18,2', 'INTERNAL'),
    ('fec6052b1fad41e99d2f97df46300bee', '2026-05-07 17:38:47', '201.170.161.229', '91294a236',  1, 'mobile_website', 'android chrome', ''),
    ('75fc523f1621495cad1fa6c4dc816359', '2026-05-07 15:11:06', '24.182.17.94',    '8087k12',    4, 'mobile_website', 'android chrome', ''),
    ('d2a882a537c54f869e8883dedaa7f280', '2026-05-07 14:55:54', '107.116.165.51',  '91400a197',  4, 'mobile_website', 'ios safari', ''),
    ('3748a2c6c895466baffa7f61a2a94042', '2026-05-07 12:28:14', '96.82.8.62',      '99362a200', 15, 'mobile_website', 'android chrome', ''),
    ('d8558080089142ce9d3a30a385c67d09', '2026-05-06 14:39:46', '72.83.150.120',   '91280a558',  1, 'mobile_website', 'ios chrome', ''),
    ('2befa8363a2244558e65ae25993f5933', '2026-05-06 15:11:16', '172.56.203.22',   '92397a114',  1, 'mobile_website', 'ios safari', ''),
    ('a67b9086609d4c71bb69ac77905d7349', '2026-05-06 17:38:07', '12.74.221.64',    '2515t106',   1, 'mobile_website', 'android chrome', ''),
    ('787cf6c976224d02a75a38f952b6de20', '2026-05-06 10:35:31', '172.59.216.243',  '91400a196',  1, 'mobile_website', 'android chrome', ''),
    ('10a7a1db851b4786bab1b733a53150ac', '2026-05-06 10:56:37', '172.56.141.47',   '91253a244',  2, 'mobile_website', 'android chrome', ''),
    ('74bbdba26b4e4aa3b05e855a3df4771a', '2026-05-05 19:39:42', '50.233.58.1',     '91124a050',  1, 'mobile_website', 'ios safari', 'INTERNAL'),
    ('81e6173b6046452fa152c78fb8bed9ab', '2026-05-05 14:33:10', '50.151.209.1',    '90185a716',  1, 'mobile_website', 'ios safari', ''),
    ('ab5b5808340c4db9b7f11a80d393f168', '2026-05-05 13:01:32', '174.227.151.85',  '5682k52',    3, 'mobile_website', 'android chrome', ''),
    ('7f8b2fde15094d04a4496a1a877cbdea', '2026-05-05 10:30:35', '66.141.84.73',    '91263a851',  1, 'mobile_website', 'ios safari', ''),
    ('3652137f45734c828cd0a32378bc716b', '2026-05-04 15:09:45', '71.95.203.178',   '91000a429',  1, 'mobile_website', 'ios safari', ''),
    ('80bdbb09935247d9b3a4dbc258b174b1', '2026-05-03 10:10:42', '27.44.125.108',   '1095k65',    1, 'mobile_website', 'android chrome', ''),
    ('343f84a50ba44d869e843f3e2b969ad2', '2026-05-10 09:54:54', '12.173.72.2',     '94820a249',  1, 'mobile_website', 'ios safari', ''),
    ('84e1a820d4b242d0a82961091c4f4043', '2026-05-06 09:43:51', '23.123.167.146',  '90764a330',  2, 'mobile_website', 'ios safari', 'FATHOM'),
    ('3fb4d3b4264840d4b3cd5ad347c52c23', '2026-05-04 11:09:37', '174.230.13.98',   '91287a268',  1, 'mobile_website', 'ios safari', ''),
    ('66320fceb99744dbba86e84a323a5fd7', '2026-05-01 10:06:25', '50.151.209.1',    '91251a242',  1, 'mobile_website', 'ios safari', ''),
    ('b379388a439749c59205461103e28238', '2026-05-01 10:06:42', '50.151.209.1',    '91251a242',  3, 'mobile_website', 'ios safari', ''),
]
print(f"Paste total rows: {len(PASTE)}")
print(f"Paste post-internal-IP-filter rows: {sum(1 for r in PASTE if r[7] != 'INTERNAL')}")
print(f"Paste unique visits (all): {len({r[0] for r in PASTE})}")
print(f"Paste unique visits (post-filter): {len({r[0] for r in PASTE if r[7] != 'INTERNAL'})}")

# Compare to current Cleaned Up tab
wb = load_workbook('/home/user/scheduling/QR_Code_Data_enriched.xlsx', data_only=True)
ws = wb['Cleaned Up']
hdr = [c.value for c in ws[1]]
vid_col = hdr.index('visit_id')
ts_col  = hdr.index('search_start_user_action_server_ts')
ip_col  = hdr.index('client_ip')

current_may = []  # (vid, ts, ip)
for row in ws.iter_rows(min_row=2, values_only=True):
    ts = row[ts_col]
    if ts and ts.month == 5:
        current_may.append((row[vid_col], ts.strftime('%Y-%m-%d %H:%M:%S'), row[ip_col]))

current_vids = {v for v, _, _ in current_may}
print(f"\nCurrent Cleaned Up May rows: {len(current_may)}; unique visits: {len(current_vids)}")

paste_vids = {r[0] for r in PASTE if r[7] != 'INTERNAL'}
new_vids = paste_vids - current_vids
print(f"\nNew visits in paste not in Cleaned Up ({len(new_vids)}):")
for r in PASTE:
    if r[0] in new_vids:
        print(f"  {r[1]}  {r[0][:12]}...  ip={r[2]}  part={r[3]}")

dropped_vids = current_vids - paste_vids
print(f"\nVisits in Cleaned Up not in paste ({len(dropped_vids)}):")
for vid, ts, ip in current_may:
    if vid in dropped_vids:
        print(f"  {ts}  {vid[:12]}...  ip={ip}")
