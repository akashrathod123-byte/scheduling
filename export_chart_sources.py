"""Export the QR weekly chart in vector formats + a clean CSV of the data."""
from openpyxl import load_workbook
from collections import Counter
from datetime import timedelta
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
import csv

XLSX = '/home/user/scheduling/QR_Code_Data_enriched.xlsx'
wb = load_workbook(XLSX, data_only=True)
ws = wb['Cleaned Up']
hdr = [c.value for c in ws[1]]
ts_col = hdr.index('search_start_user_action_server_ts')

weekly = Counter()
for row in ws.iter_rows(min_row=2, values_only=True):
    ts = row[ts_col]
    if not ts: continue
    monday = ts.date() - timedelta(days=ts.weekday())
    weekly[monday] += 1

weeks = sorted(weekly.items())

# 1. CSV of underlying data
CSV_PATH = '/home/user/scheduling/qr_weekly_data.csv'
with open(CSV_PATH, 'w', newline='') as f:
    w = csv.writer(f)
    w.writerow(['week_starting_monday', 'qr_scans'])
    for d, c in weeks:
        w.writerow([d.isoformat(), c])
print(f"CSV → {CSV_PATH}")

# 2. SVG + PDF (vector) chart
for ext in ('svg', 'pdf'):
    fig, ax = plt.subplots(figsize=(11, 4.5))
    xs = [d for d, _ in weeks]; ys = [c for _, c in weeks]
    ax.bar(xs, ys, width=5.5, color='#1f4e79', edgecolor='white')
    ax.set_xlabel('Week starting (Monday)', fontsize=10)
    ax.set_ylabel('QR scans', fontsize=10)
    ax.set_title('Weekly QR scans — Prepack items', fontsize=12, fontweight='bold', loc='left')
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    for x, y in zip(xs, ys):
        ax.text(x, y + 0.4, str(y), ha='center', va='bottom', fontsize=9)
    ax.xaxis.set_major_locator(mdates.WeekdayLocator(byweekday=mdates.MO, interval=2))
    ax.xaxis.set_major_formatter(mdates.DateFormatter('%b %d'))
    fig.autofmt_xdate()
    plt.tight_layout()
    out = f'/home/user/scheduling/qr_weekly.{ext}'
    plt.savefig(out, bbox_inches='tight')
    plt.close()
    print(f"{ext.upper()} → {out}")
