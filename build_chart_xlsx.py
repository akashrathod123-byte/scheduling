"""Build an xlsx with the weekly QR data + a native Excel bar chart."""
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import (
    RichTextProperties, Paragraph, ParagraphProperties, CharacterProperties
)
from openpyxl.styles import Font, PatternFill, Alignment
from collections import Counter
from datetime import timedelta

SRC = '/home/user/scheduling/QR_Code_Data_enriched.xlsx'
OUT = '/home/user/scheduling/qr_weekly.xlsx'

wb_src = load_workbook(SRC, data_only=True)
ws_src = wb_src['Cleaned Up']
hdr = [c.value for c in ws_src[1]]
ts_col = hdr.index('search_start_user_action_server_ts')

weekly = Counter()
for row in ws_src.iter_rows(min_row=2, values_only=True):
    ts = row[ts_col]
    if not ts: continue
    monday = ts.date() - timedelta(days=ts.weekday())
    weekly[monday] += 1
weeks = sorted(weekly.items())

wb = Workbook()
ws = wb.active
ws.title = 'Weekly QR Scans'

ws['A1'] = 'Week starting (Monday)'
ws['B1'] = 'QR scans'
for c in (ws['A1'], ws['B1']):
    c.font = Font(bold=True)
    c.fill = PatternFill('solid', fgColor='E8EEF3')

for i, (d, n) in enumerate(weeks, start=2):
    # Text label so Excel uses a category axis (wide bars) instead of date axis (thin sticks)
    ws.cell(row=i, column=1, value=d.strftime('%b %d'))
    ws.cell(row=i, column=2, value=n)

ws.column_dimensions['A'].width = 18
ws.column_dimensions['B'].width = 12

chart = BarChart()
chart.type = 'col'
chart.style = 2
chart.title = 'Weekly QR scans — Prepack items (Jan 21 – May 30, 2026)'
chart.height = 9
chart.width = 22
chart.legend = None
chart.gapWidth = 40

# Y axis: major unit of 5, no major gridlines (match doc), no axis title
chart.y_axis.majorUnit = 5
chart.y_axis.majorGridlines = None
chart.y_axis.title = None
chart.y_axis.delete = False

# X axis: show category labels (week-starting dates), no axis title
chart.x_axis.title = None
chart.x_axis.delete = False
chart.x_axis.tickLblPos = 'low'
chart.x_axis.tickLblSkip = 2  # show every other date label, like the SVG

# Rotate x-axis labels ~ -30° to mimic SVG
chart.x_axis.txPr = RichText(
    bodyPr=RichTextProperties(rot=-1800000, vert='horz'),
    p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=900)))]
)

data = Reference(ws, min_col=2, min_row=1, max_row=len(weeks) + 1, max_col=2)
cats = Reference(ws, min_col=1, min_row=2, max_row=len(weeks) + 1)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)

# Single navy series, value labels above bars (just the number, no series name/category)
series = chart.series[0]
series.graphicalProperties = GraphicalProperties(solidFill='1F4E79')
series.graphicalProperties.line.solidFill = '1F4E79'
series.dLbls = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False)

ws.add_chart(chart, 'D2')

wb.save(OUT)
print(f"Saved → {OUT}")
