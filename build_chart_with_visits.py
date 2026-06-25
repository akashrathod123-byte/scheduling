"""Add 'distinct visits' as a line over the weekly scan bars.
Outputs PNG (for the doc), SVG/PDF (for InDesign), and xlsx with native chart.
"""
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import (
    RichTextProperties, Paragraph, ParagraphProperties, CharacterProperties
)
from openpyxl.styles import Font, PatternFill
from collections import defaultdict
from datetime import timedelta
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates

SRC = '/home/user/scheduling/QR_Code_Data_enriched.xlsx'
wb_src = load_workbook(SRC, data_only=True)
ws_src = wb_src['Cleaned Up']
hdr = [c.value for c in ws_src[1]]
ts_col  = hdr.index('search_start_user_action_server_ts')
bid_col = hdr.index('browse_id')

per_week = defaultdict(lambda: {'scans': 0, 'browse': set()})
for row in ws_src.iter_rows(min_row=2, values_only=True):
    ts = row[ts_col]
    if not ts: continue
    monday = ts.date() - timedelta(days=ts.weekday())
    per_week[monday]['scans'] += 1
    if row[bid_col]:
        per_week[monday]['browse'].add(row[bid_col])

weeks = sorted(per_week)
scans  = [per_week[w]['scans'] for w in weeks]
visits = [len(per_week[w]['browse']) for w in weeks]

# ── Matplotlib renders (PNG / SVG / PDF) ─────────────────────────────────────
for ext in ('png', 'svg', 'pdf'):
    fig, ax = plt.subplots(figsize=(11, 4.8))
    bars = ax.bar(weeks, scans, width=5.5, color='#1f4e79', edgecolor='white',
                  label='QR scans')
    ax2 = ax.twinx()
    ax2.plot(weeks, visits, color='#c0504d', marker='o', linewidth=2,
             markersize=5, label='Distinct browsers')
    ax.set_ylabel('QR scans', fontsize=10)
    ax2.set_ylabel('Distinct browsers', fontsize=10, color='#c0504d')
    ax2.tick_params(axis='y', colors='#c0504d')
    ax2.set_ylim(0, max(scans) + 2)  # share scale w/ left
    ax.set_ylim(0, max(scans) + 2)
    ax.set_title('Weekly QR scans — Prepack items',
                 fontsize=12, fontweight='bold', loc='left')
    for s in ('top',): ax.spines[s].set_visible(False); ax2.spines[s].set_visible(False)
    for x, y in zip(weeks, scans):
        ax.text(x, y + 0.4, str(y), ha='center', va='bottom', fontsize=9)
    ax.xaxis.set_major_locator(mdates.WeekdayLocator(byweekday=mdates.MO, interval=2))
    ax.xaxis.set_major_formatter(mdates.DateFormatter('%b %d'))
    # combined legend
    h1, l1 = ax.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax.legend(h1 + h2, l1 + l2, loc='upper left', frameon=False, fontsize=9)
    fig.autofmt_xdate()
    plt.tight_layout()
    plt.savefig(f'/home/user/scheduling/qr_weekly_with_visits.{ext}',
                dpi=140 if ext == 'png' else None, bbox_inches='tight')
    plt.close()
    print(f"Saved → qr_weekly_with_visits.{ext}")

# ── Excel xlsx with native chart ─────────────────────────────────────────────
OUT_X = '/home/user/scheduling/qr_weekly_with_visits.xlsx'
wb = Workbook(); ws = wb.active; ws.title = 'Weekly QR Scans'
hdr_fill = PatternFill('solid', fgColor='E8EEF3')
for i, h in enumerate(['Week starting (Monday)', 'QR scans', 'Distinct browsers'], 1):
    c = ws.cell(row=1, column=i, value=h); c.font = Font(bold=True); c.fill = hdr_fill
for i, w in enumerate(weeks, 2):
    ws.cell(row=i, column=1, value=w.strftime('%b %d'))
    ws.cell(row=i, column=2, value=scans[i - 2])
    ws.cell(row=i, column=3, value=visits[i - 2])
ws.column_dimensions['A'].width = 18

chart = BarChart()
chart.type = 'col'; chart.style = 2
chart.title = 'Weekly QR scans — Prepack items (Jan 21 – May 30, 2026)'
chart.height = 9; chart.width = 22; chart.gapWidth = 40
chart.y_axis.title = None; chart.x_axis.title = None
chart.y_axis.majorUnit = 5; chart.y_axis.majorGridlines = None
chart.x_axis.tickLblSkip = 2
chart.x_axis.txPr = RichText(
    bodyPr=RichTextProperties(rot=-1800000, vert='horz'),
    p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=900)))])

data = Reference(ws, min_col=2, min_row=1, max_row=len(weeks) + 1, max_col=2)
cats = Reference(ws, min_col=1, min_row=2, max_row=len(weeks) + 1)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.series[0].graphicalProperties = GraphicalProperties(solidFill='1F4E79')
chart.series[0].dLbls = DataLabelList(showVal=True)

# Overlay line for distinct visits
line = LineChart()
ldata = Reference(ws, min_col=3, min_row=1, max_row=len(weeks) + 1, max_col=3)
line.add_data(ldata, titles_from_data=True)
line.series[0].graphicalProperties = GraphicalProperties()
line.series[0].graphicalProperties.line.solidFill = 'C0504D'
line.series[0].graphicalProperties.line.width = 22000
chart += line  # combine on same axis

ws.add_chart(chart, 'E2')
wb.save(OUT_X)
print(f"Saved → {OUT_X}")
