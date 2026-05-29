"""Build SQL to check scan→order via direct visitor_id → contact_id path."""
import re
from openpyxl import load_workbook

ENRICHED = '/home/user/scheduling/QR_Code_Data_enriched.xlsx'
OUT_SQL  = '/home/user/scheduling/check_scan_to_order_v2.sql'

ws = load_workbook(ENRICHED)['Cleaned Up']
hdr = [c.value for c in ws[1]]
idx = {h: i for i, h in enumerate(hdr)}

def cleanv(v):
    if v is None: return None
    s = str(v).strip()
    return s if s and s.upper() != 'NULL' else None

scans = []  # (visit_id, scan_date_int, part_upper)
for row in ws.iter_rows(min_row=2, values_only=True):
    vid  = cleanv(row[idx['visit_id']])
    part = cleanv(row[idx['friendly_url_txt']])
    ts   = row[idx['search_start_user_action_server_ts']]
    if not vid or not part or ts is None:
        continue
    part_upper = re.sub(r'[^A-Z0-9]', '', part.upper())
    if not part_upper: continue
    try:
        scan_date = int(ts.strftime('%Y%m%d'))
    except Exception:
        continue
    scans.append((vid, scan_date, part_upper))

print(f"Generated {len(scans):,} scan rows")

values = ",\n        ".join(
    f"('{v}', {d}, '{p}')" for (v, d, p) in scans
)

sql = f"""-- Did the scanning customer (via visitor_id -> contact_id) later order the scanned part?
WITH scans AS (
    SELECT visit_id, scan_date_int, part_upper
    FROM (VALUES
        {values}
    ) AS v(visit_id, scan_date_int, part_upper)
),
visits AS (
    SELECT visit_id, recorded_visitor_id FROM DigitalAnalytics.dbo.visit_source_202601
    UNION ALL
    SELECT visit_id, recorded_visitor_id FROM DigitalAnalytics.dbo.visit_source_202602
    UNION ALL
    SELECT visit_id, recorded_visitor_id FROM DigitalAnalytics.dbo.visit_source_202603
    UNION ALL
    SELECT visit_id, recorded_visitor_id FROM DigitalAnalytics.dbo.visit_source_202604
    UNION ALL
    SELECT visit_id, recorded_visitor_id FROM DigitalAnalytics.dbo.visit_source_202605
)
SELECT
    s.visit_id,
    s.part_upper,
    s.scan_date_int                         AS scan_date,
    v.recorded_visitor_id,
    cv.contact_id,
    os.ship_to_cmf,
    os.ship_to_cmf_name,
    COUNT(DISTINCT os.wo_id)                AS orders_after_scan,
    MIN(os.entry_date)                      AS first_order_after_scan,
    DATEDIFF(day,
             CAST(CAST(s.scan_date_int AS VARCHAR(8)) AS DATE),
             CAST(CAST(MIN(os.entry_date) AS VARCHAR(8)) AS DATE)
            )                               AS days_to_first_order,
    SUM(CAST(os.sell_value AS DECIMAL(18,2))) AS total_value_after_scan
FROM scans s
LEFT JOIN visits v                                              ON v.visit_id = s.visit_id
LEFT JOIN DigitalAnalytics.dbo.contact_visitor_source cv        ON cv.visitor_id = v.recorded_visitor_id
LEFT JOIN secure.order_source os
    ON  os.ic1_part_number = s.part_upper
    AND os.contact_id      = cv.contact_id
    AND os.entry_date     >= s.scan_date_int
    AND os.cancel_indicator = 'N'
    AND os.hk_stat_cd       = 'ACTIVE'
GROUP BY s.visit_id, s.part_upper, s.scan_date_int, v.recorded_visitor_id,
         cv.contact_id, os.ship_to_cmf, os.ship_to_cmf_name
ORDER BY orders_after_scan DESC, s.visit_id;
"""

with open(OUT_SQL, 'w') as f:
    f.write(sql)
print(f"Wrote {OUT_SQL}")
print(f"\n--- preview ---\n{sql[:1200]}")
