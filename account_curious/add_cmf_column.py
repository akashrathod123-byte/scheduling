#!/usr/bin/env python3
"""
Add column D = the actual CMF(s) for each company domain in the
"Easy Wins by Company" tab of Connecting Account Curious.xlsx.

Source of "known" people: test_with_cmf_status.xlsx, rows where
cmf_link_status == "Linked to CMF" (these have intact visitor IDs, unlike the
with_contacts CSV which Excel mangled into scientific notation). For each such
visitor we look up the CMF id in the warehouse, then list the top 15 CMFs per
domain in column D.

SPYDER USERS: edit CONFIG, press Run (F5).
Requires:  pip install pyodbc openpyxl
"""

from copy import copy
from collections import defaultdict, Counter

import openpyxl
from openpyxl.utils import get_column_letter
import pyodbc


# ===========================================================================
# EDIT THIS, THEN PRESS RUN (F5).
# ===========================================================================
CONFIG = {
    "workbook":     r"P:\Capacity Management\Management\Akash\B2B\User Models\Account Curious\Connecting Account Curious.xlsx",
    "sheet":        "Easy Wins by Company",
    # the file we already built, with clean integer IDs + status:
    "linked_file":  r"P:\Capacity Management\Management\Akash\B2B\User Models\Account Curious\test_with_cmf_status.xlsx",

    "server":   r"sqlipmds\sqlipmds",
    "database": "DigitalAnalytics",
    "driver":   "ODBC Driver 17 for SQL Server",

    "top_cmfs":  15,     # how many CMFs to list per domain
    "batch_size": 1000,
}
# ===========================================================================


MULTI_LABEL_SUFFIXES = {
    "co.uk", "org.uk", "gov.uk", "ac.uk", "com.au", "net.au", "org.au",
    "co.nz", "co.jp", "co.kr", "com.cn", "com.hk", "com.tw", "com.sg",
    "com.br", "com.mx", "co.in", "co.za", "com.tr", "co.il",
}


def registered_domain(domain):
    domain = (domain or "").strip().lower().rstrip(".")
    labels = domain.split(".")
    if len(labels) <= 2:
        return domain
    if ".".join(labels[-2:]) in MULTI_LABEL_SUFFIXES:
        return ".".join(labels[-3:])
    return ".".join(labels[-2:])


def domain_from_email(email):
    email = (email or "").strip().lower()
    if "@" not in email:
        return ""
    dom = email.rsplit("@", 1)[1].strip().strip(">").strip()
    return dom.split()[0].rstrip(".,;") if dom else ""


def find_header(ws):
    """Locate the header row and the company-domain column on the sheet."""
    for r in range(1, 12):
        for c in range(1, 8):
            val = ws.cell(r, c).value
            if val and "company" in str(val).lower():
                return r, c
    # fallback: assume header on row 3, domain in col A (matches our workbook)
    return 3, 1


def read_linked_visitors(path, target_domains):
    """From test_with_cmf_status.xlsx, yield (visitor_id, domain) for people
    who ARE linked and whose domain is one we care about."""
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    status_col = None
    kept = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            for j, h in enumerate(row):
                if h and "status" in str(h).lower():
                    status_col = j
            if status_col is None:
                status_col = 4  # our column E
            continue
        if status_col >= len(row) or row[status_col] != "Linked to CMF":
            continue
        vid = row[0]
        if isinstance(vid, str) and vid.strip().isdigit():
            vid = int(vid)
        if not isinstance(vid, int):
            continue
        comp = registered_domain(domain_from_email(row[1] if len(row) > 1 else ""))
        if comp in target_domains:
            kept += 1
            yield vid, comp
    wb.close()
    print(f"  kept {kept:,} linked people at target domains")


def fetch_cmf_ids(cur, visitor_ids):
    """visitor_id -> cmf_id (only visitors actually linked to a CMF)."""
    result = {}
    ids = list(visitor_ids)
    n = CONFIG["batch_size"]
    for i in range(0, len(ids), n):
        chunk = ids[i:i + n]
        ph = ",".join("?" * len(chunk))
        sql = f"""
            SELECT DISTINCT cv.visitor_id, cc.cmf_id
            FROM   contact_visitor_source AS cv
            JOIN   contact_cmf_source     AS cc ON cc.contact_id = cv.contact_id
            WHERE  cv.visitor_id IN ({ph}) AND cc.cmf_id IS NOT NULL
        """
        cur.execute(sql, chunk)
        for vid, cmf in cur.fetchall():
            result.setdefault(int(vid), cmf)
        print(f"  queried {min(i + n, len(ids)):,}/{len(ids):,}", end="\r")
    print()
    return result


def main():
    # ---- 1. domains from the Easy Wins tab ----
    print("Opening workbook...")
    wb = openpyxl.load_workbook(CONFIG["workbook"])
    ws = wb[CONFIG["sheet"]]
    hdr_row, dom_col = find_header(ws)
    print(f"  header on row {hdr_row}, company domain in column {get_column_letter(dom_col)}")

    domain_rows = {}  # domain -> excel row number
    for r in range(hdr_row + 1, ws.max_row + 1):
        d = ws.cell(r, dom_col).value
        if d:
            domain_rows[str(d).strip().lower()] = r
    target_domains = set(domain_rows)
    print(f"  {len(target_domains):,} company domains to resolve")

    # ---- 2. linked people at those domains (clean IDs) ----
    print("Reading linked people from test_with_cmf_status.xlsx...")
    cand = list(read_linked_visitors(CONFIG["linked_file"], target_domains))
    visitor_domain = {v: d for v, d in cand}

    # ---- 3. warehouse lookup ----
    print(f"Querying warehouse for {len(visitor_domain):,} visitors...")
    conn = pyodbc.connect(
        f"DRIVER={{{CONFIG['driver']}}};SERVER={CONFIG['server']};"
        f"DATABASE={CONFIG['database']};Trusted_Connection=yes;")
    cur = conn.cursor()
    cmf_of = fetch_cmf_ids(cur, visitor_domain.keys())
    conn.close()

    # ---- 4. domain -> Counter(cmf) ----
    domain_cmfs = defaultdict(Counter)
    for vid, cmf in cmf_of.items():
        domain_cmfs[visitor_domain[vid]][cmf] += 1

    # ---- 5. write column D ----
    out_col = 4  # column D
    ws.cell(hdr_row, out_col, f"Top {CONFIG['top_cmfs']} CMF(s)  (cmf_id: #linked people)")
    ws.cell(hdr_row, out_col).font = copy(ws.cell(hdr_row, 1).font)  # copy, not proxy

    resolved = 0
    for d, r in domain_rows.items():
        counter = domain_cmfs.get(d)
        if counter:
            resolved += 1
            top = counter.most_common(CONFIG["top_cmfs"])
            text = " | ".join(f"{cmf}: {n}" for cmf, n in top)
            extra = len(counter) - len(top)
            if extra > 0:
                text += f"  (+{extra} more)"
        else:
            text = "(no linked person found)"
        ws.cell(r, out_col, text)

    # extend the autofilter to include column D if one is set
    if ws.auto_filter.ref:
        ws.auto_filter.ref = ws.auto_filter.ref.replace(":C", ":D").replace(":c", ":D")
    ws.column_dimensions[get_column_letter(out_col)].width = 60

    wb.save(CONFIG["workbook"])
    print(f"\nDone. Resolved CMFs for {resolved:,}/{len(target_domains):,} domains.")
    print(f"Saved: {CONFIG['workbook']}")


if __name__ == "__main__":
    main()
