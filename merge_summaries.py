"""Add interaction summary columns to the enriched QR workbook."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ENRICHED  = '/home/user/scheduling/QR_Code_Data_enriched.xlsx'
SUMMARIES = '/root/.claude/uploads/329da2eb-3c16-4332-b1c4-eff6fa478b4d/f0fb8535-Interaction_Summaries.xlsx'
OUT       = '/home/user/scheduling/QR_Code_Data_enriched.xlsx'

def clean(v):
    if v is None: return None
    s = str(v).strip()
    return s if s and s.upper() != 'NULL' else None

# Build visit_id → summary
summary = {}
ws_s = load_workbook(SUMMARIES)['Summaries']
s_headers = [c.value for c in ws_s[1]]
for row in ws_s.iter_rows(min_row=2, values_only=True):
    rec = dict(zip(s_headers, row))
    vid = clean(rec.get('visit_id'))
    if vid:
        summary[vid] = rec
print(f"Summaries loaded: {len(summary):,} visits")

# Open enriched workbook, add columns
wb = load_workbook(ENRICHED)
ws = wb['Cleaned Up']
headers = [c.value for c in ws[1]]
vid_col = headers.index('visit_id') + 1
start_col = ws.max_column + 1

new_headers = ['customer_info', 'timeline', 'narrative', 'frictions']
header_font = Font(bold=True)
header_fill = PatternFill('solid', fgColor='E8EEF3')
for i, h in enumerate(new_headers):
    cell = ws.cell(row=1, column=start_col + i, value=h)
    cell.font = header_font
    cell.fill = header_fill

wrap = Alignment(wrap_text=True, vertical='top')
matched = 0
for r in range(2, ws.max_row + 1):
    vid = clean(ws.cell(row=r, column=vid_col).value)
    rec = summary.get(vid) if vid else None
    if rec:
        matched += 1
        for i, h in enumerate(new_headers):
            cell = ws.cell(row=r, column=start_col + i, value=rec.get(h))
            cell.alignment = wrap

# Column widths
widths = [40, 80, 80, 80]
for off, w in enumerate(widths):
    ws.column_dimensions[get_column_letter(start_col + off)].width = w

wb.save(OUT)
print(f"Matched {matched} of {ws.max_row - 1} rows")
print(f"Saved → {OUT}")
